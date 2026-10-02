import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# Page Setup (A4, 2cm margins)
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.79)
    section.right_margin = Inches(0.79)

# Colors
COLOR_NAVY = RGBColor(15, 42, 74)       # #0F2A4A
COLOR_BLUE = RGBColor(26, 82, 118)      # #1A5276
COLOR_GRAY = RGBColor(85, 85, 85)       # #555555
COLOR_DARK = RGBColor(30, 30, 30)       # #1E1E1E
COLOR_RED = RGBColor(180, 40, 40)       # #B42828
COLOR_GREEN = RGBColor(34, 139, 34)     # #228B22

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLUE
    return p

def add_h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK
    return p

def add_p(text, bold=False, italic=False, color=COLOR_DARK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(9.5)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return p

def add_code_block(code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, 'F4F6F7')
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    cell.width = Inches(6.69)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1A5276"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(8.0)
    run.font.color.rgb = RGBColor(20, 20, 20)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_callout(text, title='LƯU Ý QUAN TRỌNG', bg_color='EBF5FB', border_color='2E86C1'):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
    cell.width = Inches(6.69)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f'📌 {title}: ')
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(9.0)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_NAVY
    
    run_b = p.add_run(text)
    run_b.font.name = 'Arial'
    run_b.font.size = Pt(9.0)
    run_b.font.color.rgb = COLOR_DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ==================== HEADER ====================
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(2)
r1 = title_p.add_run('ĐỒ ÁN TỐT NGHIỆP CAPSTONE / RBL RESEARCH\n')
r1.font.name = 'Arial'
r1.font.size = Pt(10.5)
r1.font.bold = True
r1.font.color.rgb = COLOR_GRAY

r2 = title_p.add_run('HƯỚNG DẪN SETUP & THỰC NGHIỆM MÔ HÌNH ĐỐI SÁNH\n(BENCHMARK EXPERIMENTS EXECUTION GUIDE)')
r2.font.name = 'Arial'
r2.font.size = Pt(15)
r2.font.bold = True
r2.font.color.rgb = COLOR_NAVY

meta_p = doc.add_paragraph()
meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta_p.paragraph_format.space_after = Pt(8)
r_meta = meta_p.add_run('Dự án: ScamShield-VN | Trưởng nhóm: Trần Trung Hiếu | Ngày ban hành: 21/09/2026 | Phiên bản: 1.0')
r_meta.font.name = 'Arial'
r_meta.font.size = Pt(9.0)
r_meta.font.italic = True
r_meta.font.color.rgb = COLOR_GRAY

# Divider
div_p = doc.add_paragraph()
div_p.paragraph_format.space_after = Pt(8)
pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="0F2A4A"/></w:pBdr>')
div_p._p.get_or_add_pPr().append(pBdr)

# ==================== PHẦN I ====================
add_h1('PHẦN I: TỔNG QUAN & QUY CHUẨN THỰC NGHIỆM CHUNG')
add_p('Tài liệu này hướng dẫn chi tiết từng bước cho từng thành viên trong nhóm ScamShield-VN thực hiện chạy các mô hình AI đối chuẩn (Baselines) trên nền tảng Kaggle / Colab GPU. Mục tiêu là thu thập số liệu thực nghiệm chuẩn xác 100% để đưa vào Chương 4 (Thực nghiệm & Đánh giá) của Báo cáo Đề cương Capstone (Report 2) và Luận văn tốt nghiệp.')

add_h2('1. Ma trận Phân công Nhiệm vụ (RACI Allocation)')
tbl_assign = doc.add_table(rows=6, cols=5)
tbl_assign.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['STT', 'Thành viên', 'Mô hình phụ trách', 'Môi trường thực thi', 'Trạng thái']
col_widths = [Inches(0.5), Inches(1.3), Inches(2.2), Inches(1.5), Inches(1.19)]

for i, h in enumerate(headers):
    cell = tbl_assign.cell(0, i)
    cell.width = col_widths[i]
    set_cell_background(cell, '0F2A4A')
    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    run.font.name = 'Arial'
    run.font.size = Pt(9.0)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

data_assign = [
    ('1', 'Trần Trung Hiếu (Lead)', 'ViSoBERT + WBCE (alpha=5.0)', 'Kaggle GPU T4 x2', 'HOÀN THÀNH (Recall 98.67%)'),
    ('2', 'Hoàng Trần (Trân)', 'vinai/phobert-base-v2', 'Kaggle GPU T4 x2', 'Cần thực hiện'),
    ('3', 'Hải Phúc', 'FPTAI/vibert-base-cased', 'Kaggle GPU T4 x2', 'Cần thực hiện'),
    ('4', 'Quốc Huy', 'vinai/phobert-large', 'Kaggle GPU T4 x2', 'Cần thực hiện'),
    ('5', 'Minh Quang', 'Gemini 1.5 Flash & GPT-4o mini', 'Google AI Studio / API', 'Cần thực hiện')
]

for row_idx, data_row in enumerate(data_assign, start=1):
    bg = 'F8F9F9' if row_idx % 2 == 1 else 'FFFFFF'
    for col_idx, text_val in enumerate(data_row):
        cell = tbl_assign.cell(row_idx, col_idx)
        cell.width = col_widths[col_idx]
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        if col_idx in [0, 3, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text_val)
        run.font.name = 'Arial'
        run.font.size = Pt(8.5)
        if col_idx == 4 and 'HOÀN THÀNH' in text_val:
            run.font.bold = True
            run.font.color.rgb = COLOR_GREEN
        elif col_idx == 4:
            run.font.color.rgb = COLOR_RED

doc.add_paragraph().paragraph_format.space_after = Pt(4)

add_h2('2. Quy tắc Vàng về Tính Nhất quán Dữ liệu (Golden Rules)')
add_p('Để đảm bảo tính hợp lệ trong nghiên cứu khoa học (không bị hội đồng phản biện bắt bẻ):')
add_p('• Quy tắc 1 (Cố định phân tập): Tất cả mô hình BẮT BUỘC phải được huấn luyện và đánh giá trên cùng tập dữ liệu chuẩn đã gộp từ 2 link Hugging Face với random_state=42 (80% Train, 10% Val, 10% Test - 267 mẫu kiểm thử).')
add_p('• Quy tắc 2 (Không rò rỉ dữ liệu): Tuyệt đối không xáo trộn tập Test. Tập Test chỉ được dùng ở Cell 7 để đo điểm cuối cùng.')
add_p('• Quy tắc 3 (Trung thực khoa học): Ghi nhận đúng 100% số liệu thực tế xuất ra từ Cell 7 (Accuracy, Precision, Recall Scam, F1-Score).')

add_callout(
    'Để sử dụng GPU miễn phí 30h/tuần trên Kaggle, tài khoản cần được Verify số điện thoại (Settings -> Phone Verification). Trong giao diện Notebook, cài đặt menu bên phải: Accelerator: GPU T4 x2 và Internet: On.',
    title='CHUẨN BỊ MÔI TRƯỜNG KAGGLE'
)

# ==================== PHẦN II ====================
add_h1('PHẦN II: HƯỚNG DẪN CHI TIẾT DÀNH RIÊNG CHO TỪNG THÀNH VIÊN')

# --- MỤC 1: TRÂN ---
add_h2('1. HƯỚNG DẪN DÀNH CHO BẠN TRÂN (Chạy PhoBERT-base-v2)')
add_p('• Mục tiêu: Đo đạc độ chính xác của mô hình PhoBERT-base phiên bản 2 (mô hình chuẩn phổ biến nhất cho tiếng Việt của VinAI).')
add_p('• Quy trình thực hiện trên Kaggle Notebook:')
add_p('1. Mở Kaggle Notebook mới, chọn GPU T4 x2, bật Internet On.')
add_p('2. Chạy Cell 1 cài đặt thư viện:')
add_code_block('!pip install -q -U accelerate transformers datasets peft sentencepiece protobuf onnx onnxruntime-gpu scipy scikit-learn pyvi')
add_p('3. Khởi động lại Session (Session -> Restart Session).')
add_p('4. Chạy Cell 2 nạp dữ liệu chuẩn (Giữ nguyên code tải từ Hugging Face của nhóm).')
add_p('5. Chạy Cell 3 & Cell 4 với cấu hình mô hình:')
add_code_block('''# Cell 3 & 4: Khởi tạo PhoBERT-base-v2
import torch
from transformers import AutoTokenizer, AutoModelForSequenceClassification

MODEL_NAME = "vinai/phobert-base-v2"
tokenizer = AutoTokenizer.from_pretrained(MODEL_NAME)

# Nạp model sequence classification
model = AutoModelForSequenceClassification.from_pretrained(MODEL_NAME, num_labels=1)
print("✅ Nạp PhoBERT-base-v2 thành công!")''')
add_p('6. Chạy Cell 5 huấn luyện (Giữ nguyên batch_size=32, epochs=4, lr=2e-5).')
add_p('7. Chạy Cell 7 để lấy kết quả đánh giá trên tập Test (Accuracy, Precision, Recall Scam, F1).')
add_p('• Kết quả Trân cần gửi lại cho Hiếu: Chụp ảnh màn hình hoặc copy bảng số liệu Cell 7 + Thời gian train ở Cell 5.')

# --- MỤC 2: PHÚC ---
add_h2('2. HƯỚNG DẪN DÀNH CHO BẠN PHÚC (Chạy FPT ViBERT-base)')
add_p('• Mục tiêu: Đo đạc hiệu năng của mô hình ViBERT do FPT AI phát triển (đại diện cho trường phái kiến trúc BERT cased tiếng Việt).')
add_p('• Quy trình thực hiện trên Kaggle Notebook:')
add_p('1. Mở Kaggle Notebook mới, chọn GPU T4 x2, bật Internet On.')
add_p('2. Cài đặt thư viện (Cell 1) và Restart Session.')
add_p('3. Chạy Cell 2 nạp dữ liệu gộp.')
add_p('4. Chạy Cell 3 & Cell 4 với cấu hình mô hình ViBERT:')
add_code_block('''# Cell 3 & 4: Khởi tạo FPT ViBERT
import torch
from transformers import AutoTokenizer, AutoModelForSequenceClassification

MODEL_NAME = "FPTAI/vibert-base-cased"
tokenizer = AutoTokenizer.from_pretrained(MODEL_NAME)

model = AutoModelForSequenceClassification.from_pretrained(MODEL_NAME, num_labels=1)
print("✅ Nạp FPT ViBERT thành công!")''')
add_p('5. Chạy Cell 5 huấn luyện (epochs=4, batch_size=32).')
add_p('6. Chạy Cell 7 để lấy kết quả Test Set.')
add_p('• Kết quả Phúc cần gửi lại cho Hiếu: 4 chỉ số (Accuracy, Precision, Recall Scam, F1) và nhận xét xem ViBERT có bị nhầm ở các từ viết hoa không.')

# --- MỤC 3: HUY ---
add_h2('3. HƯỚNG DẪN DÀNH CHO BẠN HUY (Chạy PhoBERT-large)')
add_p('• Mục tiêu: Đánh giá mô hình PhoBERT-large cỡ lớn (370M tham số) để kiểm tra xem việc tăng dung lượng mô hình có giúp cải thiện Recall hay không, và độ trễ tăng bao nhiêu.')
add_callout(
    'Mô hình PhoBERT-large rất nặng. Huy BẮT BUỘC phải giảm per_device_train_batch_size xuống 16 và per_device_eval_batch_size xuống 32 ở Cell 5 để tránh lỗi Out-Of-Memory (OOM) GPU!',
    title='LƯU Ý ĐẶC BIỆT CHO BẠN HUY',
    bg_color='FDEDEC',
    border_color='C0392B'
)
add_p('• Cấu hình cụ thể cho Huy:')
add_code_block('''# Trong Cell 3 & Cell 4:
MODEL_NAME = "vinai/phobert-large"
tokenizer = AutoTokenizer.from_pretrained(MODEL_NAME)
model = AutoModelForSequenceClassification.from_pretrained(MODEL_NAME, num_labels=1)

# Trong Cell 5: SỬA LẠI BATCH SIZE
args_dict = {
    "output_dir": "./phobert_large_checkpoints",
    "eval_strategy": "epoch",
    "save_strategy": "epoch",
    "learning_rate": 1.5e-5,               # LR nhỏ hơn cho model large
    "per_device_train_batch_size": 16,     # BẮT BUỘC = 16
    "per_device_eval_batch_size": 32,      # BẮT BUỘC = 32
    "num_train_epochs": 4,
    "weight_decay": 0.01,
    "warmup_steps": 50,
    "fp16": True,
    "load_best_model_at_end": True,
    "metric_for_best_model": "recall_scam"
}''')
add_p('• Kết quả Huy cần gửi lại: Bảng điểm Cell 7 + Thời gian train (dự kiến ~5-7 phút).')

# --- MỤC 4: QUANG ---
add_h2('4. HƯỚNG DẪN DÀNH CHO BẠN QUANG (Chạy Đánh Giá Few-Shot LLMs)')
add_p('• Mục tiêu: Đánh giá độ chính xác, độ trễ và chi phí của các Large Language Models (Gemini 1.5 Flash, GPT-4o mini) ở chế độ Prompting (5-Shot In-Context Learning) mà không cần huấn luyện.')
add_p('• Hướng dẫn chạy Script Python tự động:')
add_code_block('''# Script đánh giá Few-Shot LLM trên 267 mẫu Test
import pandas as pd
import time
import requests
import json
from sklearn.metrics import accuracy_score, precision_recall_fscore_support

# 1. Tải 267 mẫu test chuẩn
URL_TEST = "https://huggingface.co/datasets/trannguyenthaituan/vietnamese_sms_dataset/raw/main/test.csv"
df_test = pd.read_csv(URL_TEST)

SYSTEM_PROMPT = """Bạn là chuyên gia an ninh mạng phân loại tin nhắn SMS tiếng Việt.
Hãy xác định tin nhắn sau là:
0: An toàn / Tin thường / OTP ngân hàng (Ham)
1: Lừa đảo / Giả mạo / Phishing / Spam độc hại (Scam)
Chỉ trả về duy nhất ký tự 0 hoặc 1."""

# Quang tích hợp gọi API Google AI Studio (Gemini Flash) hoặc OpenAI (GPT-4o mini)
# Ghi nhận danh sách dự đoán preds = [0, 1, 0, 1, ...]
# Tính toán:
# acc = accuracy_score(df_test['label'], preds)
# p, r, f1, _ = precision_recall_fscore_support(df_test['label'], preds, average='binary')
# print(f"Accuracy: {acc*100:.2f}% | Precision: {p*100:.2f}% | Recall: {r*100:.2f}% | F1: {f1*100:.2f}%")''')
add_p('• Kết quả Quang cần gửi lại: Bảng 4 chỉ số + Độ trễ trung bình mỗi request (ms) + Ước tính chi phí USD cho 267 tin nhắn.')

# ==================== PHẦN III ====================
add_h1('PHẦN III: BIỂU MẪU NỘP KẾT QUẢ & TIẾN ĐỘ THỰC HIỆN')
add_p('Tất cả các thành viên sau khi chạy xong sẽ gửi kết quả về nhóm Zalo/Discord hoặc điền trực tiếp vào bảng đối chuẩn dưới đây để Hiếu tổng hợp vào Báo cáo Đề cương:')

tbl_result = doc.add_table(rows=7, cols=8)
tbl_result.alignment = WD_TABLE_ALIGNMENT.CENTER
res_headers = ['STT', 'Mô hình', 'Phụ trách', 'Accuracy', 'Precision', 'Recall Scam', 'F1-Score', 'Độ trễ']
res_widths = [Inches(0.4), Inches(1.8), Inches(1.0), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.89)]

for i, h in enumerate(res_headers):
    cell = tbl_result.cell(0, i)
    cell.width = res_widths[i]
    set_cell_background(cell, '1A5276')
    set_cell_margins(cell, top=70, bottom=70, left=60, right=60)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    run.font.name = 'Arial'
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

data_res = [
    ('1', 'ViSoBERT + WBCE (Ours)', 'Hiếu (Lead)', '92.51%', '79.57%', '98.67%', '88.10%', '18.6 ms'),
    ('2', 'PhoBERT-base-v2', 'Trân', '... %', '... %', '... %', '... %', '... ms'),
    ('3', 'FPT ViBERT-base', 'Phúc', '... %', '... %', '... %', '... %', '... ms'),
    ('4', 'PhoBERT-large', 'Huy', '... %', '... %', '... %', '... %', '... ms'),
    ('5', 'Gemini 1.5 Flash (5-shot)', 'Quang', '... %', '... %', '... %', '... %', '... ms'),
    ('6', 'GPT-4o mini (5-shot)', 'Quang', '... %', '... %', '... %', '... %', '... ms')
]

for row_idx, row_data in enumerate(data_res, start=1):
    bg = 'EBF5FB' if row_idx == 1 else ('F8F9F9' if row_idx % 2 == 1 else 'FFFFFF')
    for col_idx, val in enumerate(row_data):
        cell = tbl_result.cell(row_idx, col_idx)
        cell.width = res_widths[col_idx]
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        p = cell.paragraphs[0]
        if col_idx in [0, 3, 4, 5, 6, 7]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(val)
        run.font.name = 'Arial'
        run.font.size = Pt(8.5)
        if row_idx == 1:
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# Footer note
p_foot = doc.add_paragraph()
p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_f = p_foot.add_run('--- HẾT TÀI LIỆU HƯỚNG DẪN ---')
r_f.font.name = 'Arial'
r_f.font.size = Pt(8.5)
r_f.font.italic = True
r_f.font.color.rgb = COLOR_GRAY

output_path = r"C:\Users\USER\RBL_ScamShield\Huong_Dan_Chay_Mo_Hinh_Doi_Sanh_ScamShield.docx"
doc.save(output_path)
print(f"✅ Successfully created document at: {output_path}")
print(f"File size: {os.path.getsize(output_path):,} bytes")
