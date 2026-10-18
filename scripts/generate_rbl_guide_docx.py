import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tc_pr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Sets subtle borders for a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def add_callout(doc, text, title="LƯU Ý QUAN TRỌNG", box_type="note"):
    """Adds a stylish callout box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = "F0F4F8" if box_type == "note" else "FFFBEB"
    border_color = "1A365D" if box_type == "note" else "D97706"
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
                          f'<w:top w:val="none"/>'
                          f'<w:right w:val="none"/>'
                          f'<w:bottom w:val="none"/>'
                          f'</w:tcBorders>')
    tc_pr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run_title = p.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(26, 54, 93) if box_type == "note" else RGBColor(180, 83, 9)
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(45, 55, 72)
    
    # spacing after callout
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def build_rbl_guide_docx(output_path):
    doc = Document()
    
    # Page Margins: 1 inch (72 pt) all sides
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(45, 55, 72) # Charcoal
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(5)

    # ---------------------------------------------------------
    # COVER / HEADER TITLE
    # ---------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sup = title_p.add_run("TÀI LIỆU HƯỚNG DẪN KỸ THUẬT RBL (RESEARCH-BASED LEARNING)\n")
    run_sup.font.size = Pt(11)
    run_sup.font.bold = True
    run_sup.font.color.rgb = RGBColor(113, 128, 150)

    run_title = title_p.add_run("SỔ TAY THIẾT LẬP OVERLEAF & QUY CHUẨN SOẠN THẢO BÀI BÁO KHOA HỌC")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(26, 54, 93) # Navy blue

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Đề tài: ScamShield-VN — Evidence-Grounded Vietnamese Scam & Phishing Detection Platform\n"
                           "Áp dụng cho: Cả 5 thành viên nhóm nghiên cứu & AI Coding Assistants (Antigravity, Gemini, Claude)")
    run_sub.font.size = Pt(10)
    run_sub.italic = True
    run_sub.font.color.rgb = RGBColor(74, 85, 104)

    # Metadata Box
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color="CBD5E0", sz="6")
    
    meta_headers = ["Chuẩn Hội Nghị", "Nguyên Tắc Dữ Liệu", "Quy Chuẩn AI Check"]
    meta_values = ["IEEEtran Conference (2-column)", "Single Source of Truth (SSOT)", "AI Score <= 18% (RBL-5b)"]
    
    for c_idx in range(3):
        cell_h = meta_table.cell(0, c_idx)
        set_cell_background(cell_h, "1A365D")
        set_cell_margins(cell_h, 80, 80, 100, 100)
        p_h = cell_h.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_h = p_h.add_run(meta_headers[c_idx])
        r_h.font.size = Pt(9.5)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(255, 255, 255)
        
        cell_v = meta_table.cell(1, c_idx)
        set_cell_background(cell_v, "F7FAFC")
        set_cell_margins(cell_v, 80, 80, 100, 100)
        p_v = cell_v.paragraphs[0]
        p_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_v = p_v.add_run(meta_values[c_idx])
        r_v.font.size = Pt(9)
        r_v.font.color.rgb = RGBColor(45, 55, 72)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper function for Section Headers
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(26, 54, 93)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(43, 108, 176)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix + " ")
            rb.bold = True
            rb.font.color.rgb = RGBColor(26, 54, 93)
        p.add_run(text)
        return p

    # ---------------------------------------------------------
    # PHẦN 1: TỔNG QUAN & QUY TRÌNH THIẾT LẬP VỚI OVERLEAF
    # ---------------------------------------------------------
    add_h1("PHẦN 1: TỔNG QUAN & QUY TRÌNH THIẾT LẬP DỰ ÁN TRÊN OVERLEAF")
    
    p = doc.add_paragraph()
    p.add_run("Tất cả các tài liệu học thuật của đề tài ScamShield-VN được quản lý tập trung trong thư mục ")
    r_code = p.add_run("paper/")
    r_code.bold = True
    p.add_run(" và đã được đóng gói sẵn thành file nén ")
    r_zip = p.add_run("paper_overleaf.zip")
    r_zip.bold = True
    p.add_run(" ở thư mục gốc của repository. Điều này giúp các thành viên có thể đưa lên Overleaf để soạn thảo trực tuyến chỉ với 1 cú click chuột mà không cần cài đặt môi trường TeX phức tạp trên máy tính cá nhân.")

    add_h2("1.1. Các bước khởi tạo dự án trên Overleaf từ file Zip")
    add_bullet("Tải file paper_overleaf.zip từ thư mục gốc của repository về máy tính.", "Bước 1 (Lấy mã nguồn):")
    add_bullet("Truy cập trang web https://www.overleaf.com, đăng nhập tài khoản của bạn.", "Bước 2 (Vào Overleaf):")
    add_bullet("Nhấn nút 'New Project' (Dự án mới) ở góc trên bên trái -> Chọn 'Upload Project' (Tải lên dự án) -> Chọn file paper_overleaf.zip vừa tải.", "Bước 3 (Tải dự án lên):")
    add_bullet("Sau khi tải lên, Overleaf sẽ tự động bung nén toàn bộ cây thư mục. Nhấn vào biểu tượng 'Menu' (góc trên bên trái giao diện chỉnh sửa):", "Bước 4 (Cấu hình Trình biên dịch):")
    
    # Sub bullets for compiler
    p_sub1 = doc.add_paragraph(style='List Bullet 2')
    p_sub1.paragraph_format.space_after = Pt(2)
    p_sub1.add_run("Compiler (Trình biên dịch): ").bold = True
    p_sub1.add_run("Chọn ")
    p_sub1.add_run("pdfLaTeX").bold = True
    p_sub1.add_run(" (hoặc XeLaTeX).")

    p_sub2 = doc.add_paragraph(style='List Bullet 2')
    p_sub2.paragraph_format.space_after = Pt(4)
    p_sub2.add_run("Main document (Tệp chính): ").bold = True
    p_sub2.add_run("Chọn ")
    p_sub2.add_run("main.tex").bold = True
    p_sub2.add_run(".")

    add_bullet("Nhấn nút xanh 'Recompile' (Biên dịch lại). Bài báo chuẩn IEEE 2 cột với bảng biểu và công thức toán học sẽ hiển thị ở khung xem trước bên phải.", "Bước 5 (Biên dịch thử):")

    add_callout(doc, 
                "Nếu Overleaf hiển thị cảnh báo font chữ hoặc gói lệnh, hãy đảm bảo trường 'Main document' được chỉ định chính xác là 'main.tex'. Không chọn các file con trong thư mục sections/ làm Main document.",
                title="MẸO OVERLEAF", box_type="note")

    add_h2("1.2. Cấu trúc thư mục mã nguồn bài báo (Thư mục paper/)")
    p = doc.add_paragraph("Mã nguồn bài báo được phân tách module hóa nghiêm ngặt để nhiều thành viên có thể viết đồng thời mà không bị xung đột (conflict):")

    files_info = [
        ("main.tex", "File master kết nối toàn bộ các sections, khai báo gói lệnh IEEEtran, metadata tác giả và cấu hình tài liệu tham khảo."),
        ("references.bib", "Cơ sở dữ liệu thư mục chứa đủ 40 tài liệu trích dẫn khoa học thu hoạch từ Systematic Literature Review (SLR)."),
        ("sections/00_abstract.tex", "Tóm tắt bài báo (150-200 từ) cấu trúc chuẩn 5 câu vàng và từ khóa Index Terms."),
        ("sections/01_intro.tex", "Section 1: Bối cảnh viễn thông, nguy cơ lừa đảo, SOTA hiện tại, 3 GAPs và 4 đóng góp cốt lõi."),
        ("sections/02_related.tex", "Section 2: Tổng hợp 40 bài báo theo 3 Theme văn hiến và Bảng Table I đối sánh."),
        ("sections/03_method.tex", "Section 3: Phương pháp luận, Dataset (N=267 Frozen), Mô hình ViSoBERT 110M, ONNX opset 18, Loss WBCE (alpha=5.0) và Protocol."),
        ("sections/04_results.tex", "Section 4: Kết quả thực nghiệm 5 runs (Table II) và kiểm định thống kê cặp McNemar (Table III)."),
        ("sections/05_discussion.tex", "Section 5: Thảo luận chuyên sâu 4 khía cạnh: Pareto-optimality, LLM Alarm fatigue, Multi-Agent Council v1.1, và Model scaling failure."),
        ("sections/06_threats.tex", "Section 6: Đánh giá 4 nhóm rủi ro (Internal, External, Construct, Conclusion) và giải pháp giảm thiểu."),
        ("sections/07_conclusion.tex", "Section 7: Tóm tắt trả lời RQ, giới hạn nghiên cứu, 3 hướng Future Work và Tuyên bố AI Usage."),
        ("claims.md", "Bảng khóa cứng 4 Luận điểm khoa học (Claim -> Evidence -> Warrant) đối soát độc lập."),
        ("quality/ai_check_log.md", "Nhật ký kiểm toán mức độ sử dụng AI viết bài (cam kết đạt chuẩn RBL-5b <= 18%)."),
        ("FULL_PAPER_PREVIEW.md", "Bản render toàn văn dạng Markdown để đọc nhanh toàn bài mà không cần mở trình biên dịch TeX.")
    ]

    for f_name, f_desc in files_info:
        add_bullet(f_desc, bold_prefix=f_name + ":")

    add_h2("1.3. Quy trình làm việc 4 bước dành cho thành viên khi nhận task")
    add_bullet("Truy cập bảng Jira nhóm, kiểm tra Task và Section bạn được phân công phụ trách. Đọc kỹ checklist và deadline của task.", "Bước 1 (Xác nhận task):")
    add_bullet("Mở file .tex tương ứng trong thư mục sections/. Toàn bộ khung nội dung, công thức toán và bảng biểu ĐÃ ĐƯỢC ĐIỀN ĐẦY ĐỦ VÀ CHÍNH XÁC.", "Bước 2 (Mở file tương ứng):")
    add_bullet("Đối soát các số liệu trong phần của bạn với 'Bản đồ số liệu đối soát (SSOT)' ở Phần 2 tài liệu này. Đọc và chỉnh sửa ngữ pháp tiếng Anh, trau chuốt câu từ cho tự nhiên và thuyết phục.", "Bước 3 (Rà soát & Chỉnh sửa):")
    add_bullet("Đây là quy định bắt buộc: Ở đầu mỗi file .tex có một khối chú thích comment '% [HƯỚNG DẪN DÀNH CHO...]'. Sau khi bạn đã kiểm tra và hoàn thiện xong phần của mình, hãy XÓA TOÀN BỘ khối comment hướng dẫn đó và các thẻ '% [TODO: ...]' đi. Sau đó recompile và nộp bài.", "Bước 4 (Xóa khối hướng dẫn):")

    # ---------------------------------------------------------
    # PHẦN 2: BẢN ĐỒ SỐ LIỆU ĐỐI SOÁT (SSOT)
    # ---------------------------------------------------------
    add_h1("PHẦN 2: BẢN ĐỒ SỐ LIỆU ĐỐI SOÁT (SINGLE SOURCE OF TRUTH - SSOT)")
    p = doc.add_paragraph("Theo Quy chế Nghiên cứu Khoa học (RESEARCH_RULES.md), ")
    p.add_run("tuyệt đối không được bịa đặt (hallucinate) số liệu").bold = True
    p.add_run(". Mọi con số xuất hiện trong bài báo IEEE bắt buộc phải bắt nguồn từ các file dữ liệu và kết quả thực nghiệm có sẵn trong repository theo bảng tra cứu dưới đây:")

    ssot_table = doc.add_table(rows=1, cols=4)
    ssot_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(ssot_table, color="CBD5E0", sz="6")

    headers = ["Thông số / Chỉ số", "Giá trị chuẩn (Bắt buộc dùng)", "File gốc trong Repository", "Section sử dụng"]
    for i, h in enumerate(headers):
        cell = ssot_table.cell(0, i)
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    ssot_data = [
        ("Tổng số tin SMS", "2.665 tin (1.918 Benign - 72.0%, 747 Scam - 28.0%)", "data/raw/vietnamese_sms_merged_all.csv\ndata/raw/README.md", "§3.1, §1"),
        ("Tập niêm phong (Frozen Test Set)", "N = 267 tin (192 Ham, 75 Scam) - Tỷ lệ 10%, seed 42", "data/raw/README.md", "Abstract, §1, §3.1, §4, §5"),
        ("Độ tin cậy gán nhãn", "Cohen's kappa = 0.88 (Near-perfect agreement)", "data/raw/README.md", "§3.1, §6.3"),
        ("ViSoBERT + WBCE Scam Recall", "95.62% ± 2.25% (Cao nhất trong các model chạy biên)", "results/summary.csv", "Abstract, §1, §4, §5, §7"),
        ("ViSoBERT + WBCE Precision", "89.79% ± 5.15%", "results/summary.csv", "Abstract, §4, §5"),
        ("ViSoBERT + WBCE Accuracy", "95.65% ± 1.54%", "results/summary.csv", "§4 (Table II)"),
        ("ViSoBERT + WBCE Macro-F1", "92.51% ± 2.39%", "results/summary.csv", "Abstract, §4, §7"),
        ("PhoBERT-large Benchmark", "Accuracy 95.95% ± 1.57%, Scam Recall 92.33% ± 3.44%", "results/summary.csv", "§4 (Table II), §5.4"),
        ("Gemini 3.5 Flash Benchmark", "Recall 100.00% nhưng Precision sụt giảm còn 64.47% ± 1.52%", "results/summary.csv", "Abstract, §1, §4, §5.2"),
        ("Kiểm định McNemar (ViSoBERT vs Gemini)", "chi2 = 30.625, p = 0.0000 (p < 0.0001), Odds Ratio = 19.00", "results/mcnemar_analysis.csv", "Abstract, §4 (Table III), §7"),
        ("Kiểm định McNemar (ViSoBERT vs Baselines)", "Odds Ratio = 2.80 - 3.00 (Vượt trội về discordance positive)", "results/mcnemar_analysis.csv", "§4 (Table III)"),
        ("Tham số mô hình & ONNX", "110M params, opset=18, size 390 MB, latency 15.8 ms", "models/scamshield_config.json\nmodels/visobert_tier1.onnx", "Abstract, §3.2"),
        ("Trọng số hàm mất mát (alpha)", "alpha = 5.0 (Phạt gấp 5 lần khi bỏ lọt tin nhắn lừa đảo)", "team-synthesis/proposal.md (§5.3)", "Abstract, §3.3, §5.1"),
        ("Ngưỡng định tuyến Tier 1 -> Tier 2", "tau = 0.90 (Xử lý tại biên nếu P >= 0.90, chuyển Hội đồng nếu < 0.90)", "team-synthesis/proposal_amendment_v1.1.md", "§1, §3.4, §5.3"),
        ("Danh sách 5 Random Seeds", "[42, 100, 123, 999, 2026]", "results/summary_5runs_detailed.csv", "§3.5, §4.1, §6.1")
    ]

    for row_idx, row_data in enumerate(ssot_data):
        row = ssot_table.add_row()
        bg = "FFFFFF" if row_idx % 2 == 0 else "F7FAFC"
        for col_idx, cell_value in enumerate(row_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(cell_value)
                r.font.size = Pt(8.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(26, 54, 93)
            elif col_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(cell_value)
                r.font.size = Pt(8.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(197, 48, 48) if "%" in cell_value or "alpha" in cell_value else RGBColor(45, 55, 72)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(cell_value)
                r.font.size = Pt(8.0)
                r.font.color.rgb = RGBColor(74, 85, 104)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------
    # PHẦN 3: WRITING CONVENTIONS (QUY CHUẨN SOẠN THẢO BÀI BÁO IEEE)
    # ---------------------------------------------------------
    add_h1("PHẦN 3: WRITING CONVENTIONS (QUY CHUẨN SOẠN THẢO BÀI BÁO IEEE)")
    
    add_h2("3.1. Phong cách và giọng văn học thuật (Academic Tone)")
    add_bullet("Sử dụng thì Quá khứ đơn (Past Simple) khi mô tả các bước thực nghiệm nhóm đã tiến hành ('We curated a unified corpus...', 'Models were trained utilizing the AdamW optimizer...').", "Quy tắc thì (Tense):")
    add_bullet("Sử dụng thì Hiện tại đơn (Present Simple) khi diễn giải các sự thật khoa học, nội dung đang hiển thị trên bảng/hình hoặc chân lý toán học ('Table II presents the empirical results...', 'Equation (1) defines the loss function...').", "Chân lý & Bảng biểu:")
    add_bullet("Tuyệt đối KHÔNG sử dụng từ viết tắt hội thoại: Viết 'do not' thay vì 'don't', viết 'cannot' thay vì 'can't', viết 'it is' thay vì 'it's'.", "Cấm viết tắt đời thường:")
    add_bullet("Tránh dùng từ ngữ cảm tính ('extremely good', 'miraculously'). Thay bằng thuật ngữ định lượng chính xác ('demonstrated statistically significant improvement (p < 0.001)', 'achieved superior Scam Recall of 95.62%').", "Khách quan định lượng:")

    add_h2("3.2. Quy chuẩn Trình bày Toán học & Ký hiệu LaTeX")
    add_bullet(r"Biến số đơn lẻ phải in nghiêng trong môi trường toán: $y_i$, $\hat{y}_i$, $N$, $\alpha$, $\tau$.", "Biến số:")
    add_bullet(r"Các hàm mất mát hoặc tên viết tắt nhiều chữ số viết bằng \text{...}: $\mathcal{L}_{\text{WBCE}}$, $\text{Recall}$, $\text{FPR}$.", "Tên hàm / Chỉ số:")
    add_bullet(r"Mọi công thức quan trọng đều phải đánh số phương trình tự động bằng \begin{equation} và gán nhãn \label{eq:...}. Khi nhắc lại trong bài báo, bắt buộc dùng \eqref{eq:...} hoặc Eq.~\eqref{eq:...}.", "Đánh số công thức:")

    add_h2("3.3. Quy chuẩn Bảng biểu (Tables) & Hình ảnh (Figures)")
    add_bullet(r"Tiêu đề của BẢNG (\caption{...}) BẮT BUỘC đặt Ở TRÊN bảng. Dùng bảng 2 cột bằng môi trường \begin{table*}[t].", "Vị trí Caption Bảng:")
    add_bullet(r"Tiêu đề của HÌNH (\caption{...}) BẮT BUỘC đặt Ở DƯỚI hình. Dùng hình 2 cột bằng môi trường \begin{figure*}[t].", "Vị trí Caption Hình:")
    add_bullet(r"Mọi bảng và hình xuất hiện trong bài báo đều PHẢI ĐƯỢC NHẮC ĐẾN ít nhất một lần trong văn bản chính (ví dụ: 'As illustrated in Fig.~\ref{fig:model_perf}...' hoặc 'Table~\ref{tab:benchmark_summary} demonstrates...'). Tuyệt đối không để bảng/hình trôi nổi mà không có lời dẫn.", "Dẫn chiếu bắt buộc:")

    add_h2("3.4. Nguyên tắc Nghiên cứu Cốt lõi (RESEARCH_RULES.md)")
    add_bullet("Data & Edge Pipeline Specialist (LR) và Metrics & Statistics Specialist (MS) là 2 vai trò độc lập. Người chuẩn bị dữ liệu và xuất mô hình không được tự chấm điểm hoặc chạy kiểm định thống kê.", "Phân lập vai trò (LR != MS):")
    add_bullet("Section 4 (Results) chỉ có nhiệm vụ báo cáo sự thật số liệu trung thực, nghiêm cấm suy diễn chủ quan hoặc giải thích nguyên nhân tại §4. Toàn bộ phần lý giải nguyên nhân, cơ chế và trade-off phải chuyển sang Section 5 (Discussion).", "Tách bạch §4 Results và §5 Discussion:")
    add_bullet("Không được thay đổi Research Question hoặc đổi metric sau khi đã nhìn thấy kết quả kiểm định để làm đẹp số liệu. Mọi đánh giá phải tuân thủ đúng cam kết tiền đăng ký.", "Cấm HARKing (Hypothesizing After the Results are Known):")
    add_bullet("Được phép dùng AI hỗ trợ trau chuốt ngữ pháp tiếng Anh, nhưng toàn bộ ý tưởng khoa học, số liệu thực nghiệm và lập luận phải do con người làm chủ. Điểm AI Detector bắt buộc duy trì <= 18% theo chuẩn RBL-5b.", "Giới hạn sử dụng AI (AI Usage Declaration):")

    # ---------------------------------------------------------
    # PHẦN 4: HỆ THỐNG FILE MARKDOWN THIẾT YẾU CHO AI AGENT
    # ---------------------------------------------------------
    add_h1("PHẦN 4: HỆ THỐNG FILE MARKDOWN THIẾT YẾU CHO AI AGENT & THÀNH VIÊN")
    p = doc.add_paragraph("Khi thành viên sử dụng AI Coding Assistants (Google Antigravity, Gemini, ChatGPT, Claude) để hỗ trợ viết hoặc review bài báo, ")
    p.add_run("bắt buộc phải chỉ định Agent đọc các file .md tham chiếu dưới đây").bold = True
    p.add_run(" để tránh trường hợp AI bịa số hoặc lạc đề:")

    md_table = doc.add_table(rows=1, cols=3)
    md_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(md_table, color="CBD5E0", sz="6")

    md_headers = ["Nhóm tài liệu", "Đường dẫn file trong Repository", "Mục đích sử dụng & Nội dung chính"]
    for i, h in enumerate(md_headers):
        cell = md_table.cell(0, i)
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    md_files = [
        ("Cẩm nang Quy chế RBL\n(rbl_docs/)", "rbl_docs/01_tong_quan.md\nrbl_docs/02_khung_danh_gia.md\nrbl_docs/03_huong_dan_slr.md\nrbl_docs/04_nghien_cuu_thuc_nghiem.md\nrbl_docs/05_huong_dan_viet.md\nrbl_docs/06_mau_tai_lieu.md", "Kim chỉ nam của hội đồng RBL: quy định mục tiêu 5 pha, thang điểm rubric, tiêu chuẩn PRISMA, thiết kế thí nghiệm, cấu trúc chuẩn từng section bài báo và mẫu khai báo AI."),
        ("Tài liệu Cốt lõi của Đề tài\n(team-synthesis/)", "team-synthesis/proposal.md\nteam-synthesis/proposal_amendment_v1.1.md\nteam-synthesis/evidence-table-merged.md\nteam-synthesis/gap-list.md", "Đề cương nghiên cứu gốc, bản cập nhật Hội đồng 5 AI giải quyết Alarm Fatigue, bảng tổng hợp trích dẫn 40 bài báo khoa học và danh sách các khoảng trống nghiên cứu đã được bảo vệ."),
        ("Bằng chứng & Luận điểm\n(paper/)", "paper/claims.md\npaper/WRITING_GUIDE_AND_ASSIGNMENT.md\npaper/quality/ai_check_log.md", "Bảng đối chiếu Claim -> Evidence -> Warrant cho 4 luận điểm khoa học, lộ trình 3 đợt viết chi tiết và log kiểm toán AI writing check."),
        ("Nhật ký & Dữ liệu Thực nghiệm\n(Repository Root)", "notes.md\ndata/raw/README.md\nresults/summary.csv\nresults/mcnemar_analysis.csv", "Nhật ký thí nghiệm ghi lại các thảo luận học thuật, giải thích chi tiết nguồn dữ liệu SMS, số liệu chi tiết 5 runs và kết quả kiểm định giả thuyết thống kê.")
    ]

    for row_idx, (g_name, g_path, g_desc) in enumerate(md_files):
        row = md_table.add_row()
        bg = "FFFFFF" if row_idx % 2 == 0 else "F7FAFC"
        for col_idx, text_val in enumerate([g_name, g_path, g_desc]):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text_val)
                r.font.size = Pt(8.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(26, 54, 93)
            elif col_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text_val)
                r.font.size = Pt(8.0)
                r.font.bold = True
                r.font.color.rgb = RGBColor(43, 108, 176)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text_val)
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(45, 55, 72)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("4.1. Prompt Mẫu Chuẩn Quốc Tế khi ra lệnh cho AI Agent viết/sửa bài")
    p = doc.add_paragraph("Để AI Agent không bịa số và tuân thủ tuyệt đối quy định học thuật, bạn hãy sao chép mẫu chỉ thị (Prompt) dưới đây khi làm việc với AI:")

    prompt_box = doc.add_table(rows=1, cols=1)
    prompt_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_p = prompt_box.cell(0, 0)
    c_p.width = Inches(6.5)
    set_cell_background(c_p, "F8FAFC")
    set_cell_margins(c_p, 120, 120, 150, 150)
    set_table_borders(prompt_box, color="CBD5E0", sz="8")
    
    p_pr = c_p.paragraphs[0]
    p_pr.paragraph_format.line_spacing = 1.15
    p_pr.paragraph_format.space_after = Pt(0)
    prompt_text = (
        "BẠN LÀ CHUYÊN GIA HỖ TRỢ VIẾT BÀI BÁO KHOA HỌC IEEE TRANSACTIONS CHO DỰ ÁN SCAMSHIELD-VN.\n\n"
        "1. NHIỆM VỤ: Hãy hỗ trợ tôi rà soát và trau chuốt nội dung file [ĐIỀN TÊN FILE, VD: paper/sections/03_method.tex].\n"
        "2. NGUYÊN TẮC BẮT BUỘC (ZERO HALLUCINATION):\n"
        "   - Bạn PHẢI đọc kỹ các file tham chiếu: rbl_docs/05_huong_dan_viet.md, paper/claims.md, results/summary.csv, và data/raw/README.md.\n"
        "   - Tuyệt đối KHÔNG ĐƯỢC TỰ BỊA ĐẶT bất kỳ con số nào. Mọi số liệu bắt buộc phải lấy chính xác từ results/summary.csv (Recall 95.62%, F1 92.51%, N=267 Frozen Test, alpha=5.0).\n"
        "   - Tuyệt đối KHÔNG chèn tên riêng cá nhân vào bài viết. Chỉ xưng hô theo danh xưng khoa học hoặc ngôi thứ nhất số nhiều trang trọng ('We').\n"
        "3. PHONG CÁCH: Văn phong học thuật tiếng Anh trang trọng, chuẩn IEEE, thì Quá khứ cho thí nghiệm đã làm, thì Hiện tại cho bảng biểu.\n"
        "4. KẾT QUẢ ĐẦU RA: Giữ nguyên các tag LaTeX, chỉ chỉnh sửa câu chữ tiếng Anh cho mượt mà, lưu ý xóa bỏ các khối hướng dẫn comment khi đã hoàn thành."
    )
    r_pr = p_pr.add_run(prompt_text)
    r_pr.font.name = "Courier New"
    r_pr.font.size = Pt(8.5)
    r_pr.font.color.rgb = RGBColor(45, 55, 72)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------
    # PHẦN 5: CHECKLIST HOÀN THÀNH TỪNG PHẦN DÀNH CHO THÀNH VIÊN
    # ---------------------------------------------------------
    add_h1("PHẦN 5: CHECKLIST HOÀN THÀNH TỪNG PHẦN DÀNH CHO THÀNH VIÊN")
    p = doc.add_paragraph("Mỗi thành viên khi đảm nhận bất kỳ Section nào cần tuân thủ bảng kiểm soát chất lượng 3 giai đoạn sau đây:")

    add_bullet("Kiểm tra mã task Jira của mình. Mở file .tex tương ứng trên Overleaf và mở file số liệu gốc tương ứng trong repo.", "Giai đoạn 1 (Chuẩn bị):")
    add_bullet("Rà soát từng câu văn tiếng Anh. Kiểm tra từng biến số toán học. Đối soát từng con số phần trăm và độ lệch chuẩn với SSOT (Phần 2).", "Giai đoạn 2 (Thực hiện):")
    add_bullet("Thực hiện đầy đủ 4 hành động trước khi đóng task:", "Giai đoạn 3 (Bàn giao & Nghiệm thu):")

    actions = [
        "1. Xóa toàn bộ khối comment hướng dẫn '% [HƯỚNG DẪN DÀNH CHO...]' ở đầu mục của mình.",
        "2. Xóa toàn bộ các thẻ '% [TODO: ...]' còn sót lại trong mục.",
        "3. Bấm 'Recompile' trên Overleaf kiểm tra lỗi biên dịch (0 Errors, 0 Warnings nghiêm trọng).",
        "4. Tải file .tex đã sửa về hoặc commit lên Git, sau đó đánh dấu 'Done' trên Jira."
    ]
    for act in actions:
        p_act = doc.add_paragraph(style='List Bullet 2')
        p_act.paragraph_format.space_after = Pt(2)
        r_act = p_act.add_run(act)
        r_act.font.bold = True
        r_act.font.color.rgb = RGBColor(26, 54, 93)

    add_callout(doc, 
                "Tuân thủ nghiêm ngặt quy trình trên sẽ đảm bảo bài báo của nhóm đạt chất lượng công bố quốc tế, bảo vệ thành công trước Hội đồng RBL với điểm số tối đa!", 
                title="KẾT LUẬN", box_type="note")

    # Save document
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\Users\USER\RBL_ScamShield\docs"
    out_file = os.path.join(out_dir, "HUONG_DAN_SETUP_VA_LAM_TASK_RBL_OVERLEAF.docx")
    build_rbl_guide_docx(out_file)
